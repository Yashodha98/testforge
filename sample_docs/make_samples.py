"""Builds SYNTHETIC Confluence-style requirement exports used for demos/tests.

Content is invented for this PoC (modelled on a hospital management system);
it contains no real client requirements or patient data.
Run:  python sample_docs/make_samples.py
"""
from pathlib import Path

import docx

HERE = Path(__file__).parent
PAGE = "HMS Release 4.2 - Outpatient Module - Requirements"
INTRO = ("This page describes the outpatient registration and appointment features for the Hospital Management "
         "System (HMS) release 4.2. Users: front-desk receptionists and patients using the mobile app. "
         "Owner: Product Owner, Outpatient. Status: Approved for Sprint 18.")
STORIES = [
    ("US-101 Register a new patient",
     "As a front-desk receptionist, I want to register a new patient, so that the patient can be booked for consultations and billed correctly.",
     [("AC1", "Given the receptionist is on the New Patient screen, when all mandatory fields are completed and Save is clicked, then a unique Medical Record Number (MRN) is generated and shown within 2 seconds."),
      ("AC2", "The patient full name must be between 2 and 60 characters."),
      ("AC3", "Age is calculated from the date of birth and must be between 0 and 120 years; a future date of birth is rejected."),
      ("AC4", "The National ID / Iqama number must be exactly 10 digits and start with 1 or 2."),
      ("AC5", "When a patient with the same National ID already exists, the message \"Patient already registered\" is shown and the existing MRN is displayed.")]),
    ("US-102 Book an outpatient appointment",
     "As a patient using the hospital mobile app, I want to book an appointment with a doctor, so that I can be seen without calling the hospital.",
     [("AC1", "Available 15-minute slots for the selected doctor and date are displayed, and booked slots cannot be selected."),
      ("AC2", "Appointments can be booked up to 30 days in advance; past dates cannot be selected."),
      ("AC3", "A patient can have a maximum of 3 active appointments; booking a 4th shows \"You already have 3 active appointments\"."),
      ("AC4", "After booking, a confirmation with doctor, date, time and booking reference is displayed and an SMS is sent.")]),
    ("US-103 Cancel an appointment",
     "As a patient, I want to cancel my appointment in the app, so that the slot can be given to another patient.",
     [("AC1", "Cancellation is allowed only at least 2 hours before the appointment time."),
      ("AC2", "Within 2 hours of the appointment the message \"Please contact the clinic to cancel\" is shown and the Cancel button is disabled."),
      ("AC3", "After cancellation the slot becomes available again immediately and a cancellation SMS is sent.")]),
]


def make_docx():
    d = docx.Document()
    d.add_heading(PAGE, 0)
    d.add_paragraph(INTRO)
    for title, story, acs in STORIES:
        d.add_heading(title, 1)
        d.add_paragraph(story)
        d.add_paragraph("Acceptance Criteria")
        t = d.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "ID", "Acceptance Criteria"
        for ac, text in acs:
            r = t.add_row().cells
            r[0].text, r[1].text = ac, text
    d.save(HERE / "Confluence_HMS_Outpatient_Requirements.docx")


def make_html():
    parts = ["<html><head><title>{0}</title><style>body{{font-family:Arial}}</style></head><body>".format(PAGE),
             "<h1>{}</h1><p>{}</p>".format(PAGE, INTRO)]
    for title, story, acs in STORIES:
        parts.append("<h2>{}</h2><p>{}</p><h3>Acceptance Criteria</h3><ul>".format(title, story))
        parts += ["<li>{}</li>".format(t.replace('"', "&quot;")) for _, t in acs]
        parts.append("</ul>")
    parts.append("</body></html>")
    (HERE / "Confluence_HMS_Outpatient_Requirements.html").write_text("\n".join(parts), encoding="utf-8")


if __name__ == "__main__":
    make_docx()
    make_html()
    print("Synthetic Confluence exports written to", HERE)
